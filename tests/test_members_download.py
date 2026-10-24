import io
from datetime import timedelta
from zoneinfo import ZoneInfo

import pytest
from django.urls import reverse
from django.utils.timezone import now
from django_scopes import scope
from eventyay.base.models import Team
from openpyxl import load_workbook

from teamshifts.exporters import MembersExporter, format_answer, format_hours
from teamshifts.models import (
    ApplicationStatus,
    CallForTeamMembers,
    QuestionVariant,
    Shift,
    ShiftAssignment,
    TeamApplicationAnswer,
    TeamApplicationQuestion,
    TeamApplicationQuestionOption,
    TeamMemberApplication,
    TeamRole,
)


@pytest.fixture(autouse=True)
def site_url(settings):
    settings.SITE_URL = "https://testserver"


@pytest.fixture
def cfm(event):
    with scope(event=event):
        return CallForTeamMembers.objects.create(event=event, active=True)


@pytest.fixture
def orga_user(event, user):
    with scope(event=event):
        team = Team.objects.create(organizer=event.organizer, name="Orga", can_change_event_settings=True, all_events=True)
        team.members.add(user)
    return user


@pytest.fixture
def lead_user(event, django_user_model):
    lead = django_user_model.objects.create_user(email="lead@example.com", password="secret")
    with scope(event=event):
        team = Team.objects.create(organizer=event.organizer, name="Leads", all_events=True, teamshifts_role="lead")
        team.members.add(lead)
    return lead


def _url(event):
    return reverse("plugins:teamshifts:members_download", kwargs={"organizer": event.organizer.slug, "event": event.slug})


def _member(event, django_user_model, email, fullname="", **kwargs):
    applicant = django_user_model.objects.create_user(email=email, password="x", fullname=fullname)
    return TeamMemberApplication.objects.create(event=event, user=applicant, status=ApplicationStatus.ACCEPTED, **kwargs)


def _rows(response):
    workbook = load_workbook(io.BytesIO(response.content))
    return [[cell.value for cell in row] for row in workbook.active.iter_rows()]


def test_format_answer_boolean_and_multiple():
    boolean = TeamApplicationQuestion(variant=QuestionVariant.BOOLEAN)
    multiple = TeamApplicationQuestion(variant=QuestionVariant.MULTIPLE)
    tz = ZoneInfo("UTC")
    assert format_answer(boolean, "true", tz) == "Yes"
    assert format_answer(boolean, "false", tz) == "No"
    assert format_answer(multiple, "S\nM\nL", tz) == "S, M, L"
    assert format_answer(multiple, "", tz) == ""


def test_format_answer_datetime_uses_event_timezone():
    question = TeamApplicationQuestion(variant=QuestionVariant.DATETIME)
    assert format_answer(question, "2026-10-03T10:00:00+00:00", ZoneInfo("Asia/Kolkata")) == "2026-10-03 15:30"
    assert format_answer(question, "2026-10-03 10:00", ZoneInfo("Asia/Kolkata")) == "2026-10-03 10:00"


def test_format_hours():
    assert format_hours(None) == 0
    assert format_hours(timedelta(hours=2, minutes=30)) == 2.5


@pytest.mark.django_db
def test_csv_escapes_formulas(client, event, cfm, orga_user, django_user_model):
    with scope(event=event):
        _member(event, django_user_model, "a@example.com", "=cmd()")
    client.force_login(orga_user)
    body = client.get(_url(event), {"_format": "default"}).content.decode("utf-8")
    assert '"=cmd()"' not in body
    assert "'=cmd()" in body


@pytest.mark.django_db
def test_xlsx_stores_formula_like_text_as_string(client, event, cfm, orga_user, django_user_model):
    with scope(event=event):
        _member(event, django_user_model, "a@example.com", "=1+1")
    client.force_login(orga_user)
    workbook = load_workbook(io.BytesIO(client.get(_url(event), {"_format": "xlsx"}).content))
    cell = next(row for row in workbook.active.iter_rows(min_row=2))[0]
    assert cell.data_type == "s"
    assert cell.value == "'=1+1"


@pytest.mark.django_db
def test_exporter_filename_uses_event_slug_and_date(event):
    assert MembersExporter(event, []).get_filename().startswith(f"{event.slug}-members-")


@pytest.mark.django_db
def test_download_xlsx_contains_members_and_answers(client, event, cfm, orga_user, django_user_model):
    with scope(event=event):
        role = TeamRole.objects.create(event=event, name="Volunteer")
        tshirt = TeamApplicationQuestion.objects.create(event=event, question="T-shirt", variant=QuestionVariant.MULTIPLE)
        TeamApplicationQuestionOption.objects.create(question=tshirt, answer="S", position=0)
        TeamApplicationQuestionOption.objects.create(question=tshirt, answer="M", position=1)
        agree = TeamApplicationQuestion.objects.create(event=event, question="Agree", variant=QuestionVariant.BOOLEAN)
        member = _member(event, django_user_model, "a@example.com", "Alice", phone="123", arrived=True, added_by_organizer=True)
        _member(event, django_user_model, "b@example.com", "Bob")
        TeamApplicationAnswer.objects.create(application=member, question=tshirt, answer="S\nM")
        TeamApplicationAnswer.objects.create(application=member, question=agree, answer="true")
        start = now()
        shift = Shift.objects.create(event=event, name="Desk", start_time=start, end_time=start + timedelta(hours=2))
        ShiftAssignment.objects.create(shift=shift, team_member=member.user, role=role)

    client.force_login(orga_user)
    response = client.get(_url(event), {"_format": "xlsx"})
    assert response.status_code == 200
    assert response["Content-Disposition"].startswith(f'attachment; filename="{event.slug}-members-')
    assert response["Content-Disposition"].endswith('.xlsx"')
    rows = _rows(response)
    header = rows[0]
    assert header[:2] == ["Name", "Email"]
    assert header[-2:] == ["T-shirt", "Agree"]
    alice = next(row for row in rows[1:] if row[1] == "a@example.com")
    assert alice[0] == "Alice"
    assert alice[2] == "123"
    assert alice[3] == "Volunteer"
    assert alice[4] == 1
    assert alice[5] == 2
    assert alice[6] == "Yes"
    assert alice[10] == "Yes"
    assert alice[-2:] == ["S, M", "Yes"]
    bob = next(row for row in rows[1:] if row[1] == "b@example.com")
    assert bob[-2:] == [None, None]
    assert bob[10] == "No"


@pytest.mark.django_db
def test_download_csv_respects_search_filter(client, event, cfm, orga_user, django_user_model):
    with scope(event=event):
        _member(event, django_user_model, "a@example.com", "Alice")
        _member(event, django_user_model, "b@example.com", "Bob")
    client.force_login(orga_user)
    response = client.get(_url(event), {"_format": "default", "q": "alice"})
    assert response.status_code == 200
    assert response["Content-Type"].startswith("text/csv")
    body = response.content.decode("utf-8-sig")
    assert "a@example.com" in body
    assert "b@example.com" not in body


@pytest.mark.django_db
def test_download_role_filter_limits_members(client, event, cfm, orga_user, django_user_model):
    with scope(event=event):
        role = TeamRole.objects.create(event=event, name="Volunteer")
        member = _member(event, django_user_model, "a@example.com", "Alice")
        _member(event, django_user_model, "b@example.com", "Bob")
        start = now()
        shift = Shift.objects.create(event=event, name="Desk", start_time=start, end_time=start + timedelta(hours=1))
        ShiftAssignment.objects.create(shift=shift, team_member=member.user, role=role)
    client.force_login(orga_user)
    response = client.get(_url(event), {"_format": "default", "role": role.pk})
    body = response.content.decode("utf-8-sig")
    assert "a@example.com" in body
    assert "b@example.com" not in body


@pytest.mark.django_db
@pytest.mark.parametrize("export_format", ["pdf", "csv-excel", "semicolon", "csv"])
def test_download_rejects_unsupported_format(client, event, cfm, orga_user, export_format):
    client.force_login(orga_user)
    assert client.get(_url(event), {"_format": export_format}).status_code == 400


@pytest.mark.django_db
def test_download_forbidden_for_team_lead(client, event, cfm, lead_user):
    client.force_login(lead_user)
    assert client.get(_url(event), {"_format": "default"}).status_code == 403


@pytest.mark.django_db
def test_download_requires_login(client, event, cfm):
    assert client.get(_url(event), {"_format": "default"}).status_code in (302, 403)


@pytest.mark.django_db
def test_members_page_shows_download_only_to_coordinators(client, event, cfm, orga_user, lead_user):
    members_url = reverse("plugins:teamshifts:members", kwargs={"organizer": event.organizer.slug, "event": event.slug})
    client.force_login(orga_user)
    assert "_format=xlsx" in client.get(members_url).content.decode()
    client.force_login(lead_user)
    assert "_format=xlsx" not in client.get(members_url).content.decode()


@pytest.mark.django_db
@pytest.mark.parametrize("role", ["99999999999999999999", "\u00b2", "-1", "abc"])
def test_download_ignores_invalid_role_filter(client, event, cfm, orga_user, django_user_model, role):
    with scope(event=event):
        _member(event, django_user_model, "a@example.com", "Alice")
    client.force_login(orga_user)
    response = client.get(_url(event), {"_format": "default", "role": role})
    assert response.status_code == 200
    assert "a@example.com" in response.content.decode("utf-8-sig")


@pytest.mark.django_db
@pytest.mark.parametrize("role", ["99999999999999999999", "\u00b2"])
def test_members_page_ignores_invalid_role_filter(client, event, cfm, orga_user, role):
    members_url = reverse("plugins:teamshifts:members", kwargs={"organizer": event.organizer.slug, "event": event.slug})
    client.force_login(orga_user)
    assert client.get(members_url, {"role": role}).status_code == 200


@pytest.mark.django_db
def test_deleted_question_is_removed_from_export_and_answers_are_deleted(client, event, cfm, orga_user, django_user_model):
    with scope(event=event):
        question = TeamApplicationQuestion.objects.create(event=event, question="Diet", variant=QuestionVariant.STRING)
        answered = _member(event, django_user_model, "a@example.com", "Alice")
        TeamApplicationAnswer.objects.create(application=answered, question=question, answer="Vegan")
    client.force_login(orga_user)
    delete_url = reverse("plugins:teamshifts:question_delete", kwargs={"organizer": event.organizer.slug, "event": event.slug, "pk": question.pk})
    assert client.post(delete_url).status_code == 302
    with scope(event=event):
        assert not TeamApplicationQuestion.objects.filter(pk=question.pk).exists()
        assert not TeamApplicationAnswer.objects.filter(application=answered).exists()

    rows = _rows(client.get(_url(event), {"_format": "xlsx"}))
    assert "Diet" not in rows[0]
    assert all("Vegan" not in row for row in rows)


@pytest.mark.django_db
def test_inactive_question_is_not_exported(client, event, cfm, orga_user, django_user_model):
    with scope(event=event):
        question = TeamApplicationQuestion.objects.create(event=event, question="Diet", variant=QuestionVariant.STRING, active=False)
        answered = _member(event, django_user_model, "a@example.com", "Alice")
        TeamApplicationAnswer.objects.create(application=answered, question=question, answer="Vegan")
    client.force_login(orga_user)
    rows = _rows(client.get(_url(event), {"_format": "xlsx"}))
    assert "Diet" not in rows[0]
    assert all("Vegan" not in row for row in rows)
    with scope(event=event):
        assert TeamApplicationAnswer.objects.filter(application=answered, question=question).exists()


@pytest.mark.django_db
def test_question_delete_page_warns_answers_are_permanently_deleted(client, event, cfm, orga_user):
    with scope(event=event):
        question = TeamApplicationQuestion.objects.create(event=event, question="Diet", variant=QuestionVariant.STRING)
    client.force_login(orga_user)
    delete_url = reverse("plugins:teamshifts:question_delete", kwargs={"organizer": event.organizer.slug, "event": event.slug, "pk": question.pk})
    body = client.get(delete_url).content.decode()
    assert "Answers are permanently deleted." in body
    assert "Deactivate the question instead to stop asking it while keeping its answers." in body
